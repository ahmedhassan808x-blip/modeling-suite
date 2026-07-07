"""
Presentation & export layer for the merger model.

    python3 -m merger.report MSFT_AAPL_merger.xlsx
"""

import re
import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook  # noqa: E402

from merger.model_builder import (  # noqa: E402
    CHECKS_PASS_CELL, CO, DEAL, PF,
)
from shared import charts  # noqa: E402
from shared.exports.docx_memo import Memo  # noqa: E402
from shared.exports.pdf import to_pdf  # noqa: E402
from shared.exports.pptx_deck import Deck  # noqa: E402
from shared.recalc import recalculate  # noqa: E402

FC = list("FGHIJ")

LIMITS = [
    "Cost synergies only in the base case; revenue synergies excluded by "
    "discipline (assume them and any deal 'works'). The synergy seed is a "
    "placeholder heuristic (10% of target EBITDA) — the number is the "
    "analyst's judgment, not the model's.",
    "New acquisition debt carried flat — no paydown or refinancing path "
    "modeled; incremental financing cost is therefore mildly conservative "
    "in later years.",
    "Purchase price allocation is 'basics': straight-line intangible "
    "amortization, no deferred-tax liabilities on write-ups, no PP&E "
    "step-up, no fair-value adjustments.",
    "Compact earnings forecasts per company (no full balance-sheet / "
    "revolver build); net interest & other held flat at last actual.",
    "EPS accretion is a screening lens, not value creation — a deal can be "
    "accretive and still destroy value (and vice versa).",
]


def _probes():
    p = [CHECKS_PASS_CELL] + [f"Deal!B{r}" for r in DEAL.values()]
    for c in FC:
        p += [f"ProForma!{c}{PF[k]}" for k in
              ("acq_ni", "tgt_ni", "syn", "int_", "amort", "adj_at", "pf_ni",
               "pf_eps", "acq_eps", "acc_pct", "breakeven")]
    for i in range(5):
        r = PF["sens_first"] + i
        p += [f"ProForma!B{r}"] + [f"ProForma!{sc}{r}" for sc in "CDEFG"]
    p += [f"ProForma!{sc}{PF['sens_hdr']}" for sc in "CDEFG"]
    return p


def extract(xlsx) -> dict:
    res = recalculate(xlsx, probe_cells=_probes())
    if not res.ok:
        raise RuntimeError(f"{xlsx}: formula errors — refusing to report.\n"
                           + res.summary())
    v = res.values
    if v[CHECKS_PASS_CELL] != "PASS":
        raise RuntimeError(f"{xlsx}: M&A checks FAIL — refusing to report.")
    wb = load_workbook(xlsx)
    m = re.match(r"^Deal: (\w[\w.\-]*) acquires (\w[\w.\-]*)",
                 wb["Deal"]["A1"].value or "")
    acq, tgt = (m.group(1), m.group(2)) if m else ("?", "?")
    deal = {k: v[f"Deal!B{r}"] for k, r in DEAL.items()}
    pf = {k: [v[f"ProForma!{c}{PF[k]}"] for c in FC] for k in
          ("acq_ni", "tgt_ni", "syn", "int_", "amort", "adj_at", "pf_ni",
           "pf_eps", "acq_eps", "acc_pct", "breakeven")}
    sens = dict(
        prem=[v[f"ProForma!{sc}{PF['sens_hdr']}"] for sc in "CDEFG"],
        syn=[v[f"ProForma!B{PF['sens_first'] + i}"] for i in range(5)],
        grid=[[v[f"ProForma!{sc}{PF['sens_first'] + i}"] for sc in "CDEFG"]
              for i in range(5)])
    return dict(acq=acq, tgt=tgt, deal=deal, pf=pf, sens=sens)


def _deal_table(deal):
    return [["Deal terms", ""],
            ["Offer per share", f"${deal['offer_ps']:,.2f} "
             f"({deal['premium']:.0%} premium)"],
            ["Offer value (equity)", f"${deal['offer_val']:,.0f}mm"],
            ["Exchange ratio (stock portion)", f"{deal['xratio']:.3f}x"],
            ["Mix stock / cash / debt",
             f"{deal['mix_stock']:.0%} / {deal['mix_cash']:.0%} / "
             f"{deal['mix_debt']:.0%}"],
            ["New shares issued", f"{deal['new_shares']:,.1f}mm"],
            ["New debt (incl. fees)", f"${deal['debt_cons']:,.0f}mm"],
            ["Goodwill / intangibles",
             f"${deal['goodwill']:,.0f}mm / ${deal['intang']:,.0f}mm"],
            ["Run-rate synergies (pre-tax)", f"${deal['syn']:,.0f}mm"]]


def _sens_table(sens):
    rows = [["Synergies \\ premium"] + [f"{p:.0%}" for p in sens["prem"]]]
    for s, grid in zip(sens["syn"], sens["grid"]):
        rows.append([f"${s:,.0f}mm"] + [f"{x:+.1%}" for x in grid])
    return rows


def build_reports(xlsx, outdir=None):
    xlsx = Path(xlsx)
    outdir = Path(outdir) if outdir else xlsx.parent
    d = extract(xlsx)
    deal, pf = d["deal"], d["pf"]

    tmp = Path(tempfile.mkdtemp())
    years = [f"Year {t}" for t in range(1, 6)]
    png_acc = charts.pct_bars(years, pf["acc_pct"], tmp / "acc.png")
    adj_pre = pf["syn"][0] + pf["int_"][0] + pf["amort"][0]
    png_bridge = charts.waterfall(
        [("Acquirer NI", pf["acq_ni"][0], "total"),
         ("Target NI", pf["tgt_ni"][0], "delta"),
         ("Synergies", pf["syn"][0], "delta"),
         ("Financing", pf["int_"][0], "delta"),
         ("Intang. amort", pf["amort"][0], "delta"),
         ("Tax effect", pf["adj_at"][0] - adj_pre, "delta"),
         ("Pro forma NI", pf["pf_ni"][0], "total")],
        tmp / "bridge.png", title="Year 1 net income bridge ($mm)")

    acc1, acc5 = pf["acc_pct"][0], pf["acc_pct"][-1]
    summary = [
        f"{d['acq']} acquires {d['tgt']} at {deal['premium']:.0%} premium "
        f"(${deal['offer_ps']:,.2f}/sh, ${deal['offer_val']:,.0f}mm equity "
        f"value), {deal['mix_stock']:.0%} stock / {deal['mix_cash']:.0%} "
        f"cash / {deal['mix_debt']:.0%} debt.",
        f"Year 1 {'accretion' if acc1 >= 0 else 'dilution'} of {acc1:+.1%}, "
        f"reaching {acc5:+.1%} by Year 5 as synergies phase in "
        "(50/75/100%).",
        f"Run-rate cost synergies ${deal['syn']:,.0f}mm pre-tax "
        f"(placeholder seed); breakeven synergies for EPS-neutral Year 1: "
        f"${pf['breakeven'][0]:,.0f}mm.",
        f"PPA: ${deal['intang']:,.0f}mm intangibles amortized over "
        f"{deal['amort_years']:.0f} years, ${deal['goodwill']:,.0f}mm "
        "goodwill.",
        "All integrity checks PASS (mix, sources=uses, goodwill ≥ 0, "
        "independent Y1 re-derivation) — verified by recalculation.",
    ]

    footer = (f"{d['acq']} / {d['tgt']} — merger consequences | "
              "modeling-suite | $mm | verify assumptions before use")
    deck = Deck(footer)
    deck.title_slide(f"{d['acq']} acquires {d['tgt']}",
                     "Merger consequences — accretion / dilution",
                     "Live Excel model: blue = input, black = formula, "
                     "green = cross-sheet link")
    s = deck.content_slide("Executive summary")
    deck.add_bullets(s, summary, 0.5, 1.3, 6.6, 5.4, size=12)
    deck.add_table(s, _deal_table(deal), 7.5, 1.5, 5.3, first_col_w=2.7)

    s = deck.content_slide("Accretion / (dilution) walk")
    deck.add_image(s, png_acc, 0.5, 1.5, 6.3)
    deck.add_image(s, png_bridge, 7.0, 1.5, 6.0)

    s = deck.content_slide("Sensitivity — Year 1 accretion: synergies × "
                           "premium")
    deck.add_table(s, _sens_table(d["sens"]), 0.5, 1.5, 8.4, first_col_w=2.0)
    deck.add_bullets(s, [
        f"Center cell = base case ({pf['acc_pct'][0]:+.1%}) — independently "
        "re-derived and check-verified against the ProForma sheet.",
    ], 0.5, 4.6, 11.5, 1.5, size=10)

    s = deck.content_slide("Risks & limitations")
    deck.add_bullets(s, LIMITS, 0.5, 1.4, 12.0, 5.4, size=12)
    pptx_path = deck.save(outdir / f"{xlsx.stem}_deck.pptx")

    memo = Memo(f"{d['acq']} acquires {d['tgt']} — Merger Consequences Memo",
                "Accretion/dilution analysis | modeling-suite")
    memo.heading("Executive summary")
    memo.bullets(summary)
    memo.table(_deal_table(deal))
    memo.heading("Accretion / (dilution)")
    memo.image(png_acc)
    memo.image(png_bridge)
    memo.heading("Sensitivity (synergies × premium, Year 1)")
    memo.table(_sens_table(d["sens"]))
    memo.heading("Risks & limitations")
    memo.bullets(LIMITS)
    memo.para("Generated by modeling-suite from the recalculated workbook — "
              "figures cannot diverge from the model.", small=True)
    docx_path = memo.save(outdir / f"{xlsx.stem}_memo.docx")

    pdf_path = to_pdf(pptx_path, outdir)
    shutil.rmtree(tmp, ignore_errors=True)
    return dict(pptx=pptx_path, docx=docx_path, pdf=pdf_path)


def main():
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--outdir")
    args = ap.parse_args()
    for k, p in build_reports(args.workbook, args.outdir).items():
        print(f"{k}: {p}")


if __name__ == "__main__":
    main()
